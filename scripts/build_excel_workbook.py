import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

cleaned_path = os.path.join('data', 'cleaned_data.csv')
excel_path = os.path.join('excel', 'ecommerce_analysis.xlsx')

os.makedirs('excel', exist_ok=True)

df = pd.read_csv(cleaned_path)

wb = openpyxl.Workbook()

# Style definitions
font_family = "Segoe UI"
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Navy
header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
title_font = Font(name=font_family, size=16, bold=True, color="1F4E78")
subtitle_font = Font(name=font_family, size=11, italic=True, color="595959")
kpi_label_font = Font(name=font_family, size=9, bold=True, color="595959")
kpi_val_font = Font(name=font_family, size=18, bold=True, color="1F4E78")
section_font = Font(name=font_family, size=13, bold=True, color="1F4E78")
bold_font = Font(name=font_family, size=10, bold=True)
regular_font = Font(name=font_family, size=10)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

total_border = Border(
    top=Side(style='thin', color='000000'),
    bottom=Side(style='double', color='000000')
)

kpi_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

# ---------------------------------------------------------------------------
# SHEET 1: EXECUTIVE SUMMARY
# ---------------------------------------------------------------------------
ws_summary = wb.active
ws_summary.title = "Executive Summary"
ws_summary.views.sheetView[0].showGridLines = True

# Title Block
ws_summary["B2"] = "E-COMMERCE EXECUTIVE PERFORMANCE DASHBOARD"
ws_summary["B2"].font = title_font
ws_summary["B3"] = "Financial KPI Overview, Category Performance & Strategic Highlights"
ws_summary["B3"].font = subtitle_font

# KPIs Calculations
tot_sales = df['Sales'].sum()
tot_profit = df['Profit'].sum()
tot_orders = df['Order_ID'].nunique()
tot_cust = df['Customer_ID'].nunique()
aov = tot_sales / tot_orders
margin_pct = (tot_profit / tot_sales)

kpis = [
    ("TOTAL REVENUE", f"${tot_sales:,.2f}", "B5", "C6"),
    ("TOTAL PROFIT", f"${tot_profit:,.2f}", "D5", "E6"),
    ("TOTAL ORDERS", f"{tot_orders:,}", "F5", "G6"),
    ("CUSTOMERS", f"{tot_cust:,}", "H5", "I6"),
    ("AVG ORDER VALUE", f"${aov:,.2f}", "J5", "K6"),
    ("PROFIT MARGIN", f"{margin_pct * 100:.2f}%", "L5", "M6"),
]

for label, val, top_left, bot_right in kpis:
    c_label = ws_summary[top_left]
    c_label.value = label
    c_label.font = kpi_label_font
    c_label.alignment = Alignment(horizontal="center", vertical="center")
    c_label.fill = kpi_fill
    
    # Extract row and col from bot_right
    col_letter = bot_right[0]
    row_num = int(bot_right[1:])
    c_val = ws_summary[f"{top_left[0]}{int(top_left[1])+1}"]
    c_val.value = val
    c_val.font = kpi_val_font
    c_val.alignment = Alignment(horizontal="center", vertical="center")
    c_val.fill = kpi_fill

# Category Summary Table on Executive Sheet
ws_summary["B9"] = "Category Financial Performance Summary"
ws_summary["B9"].font = section_font

cat_headers = ["Category", "Total Orders", "Units Sold", "Total Revenue ($)", "Total Profit ($)", "Profit Margin (%)"]
for col_idx, h in enumerate(cat_headers, start=2):
    cell = ws_summary.cell(row=10, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left", vertical="center")

cat_df = df.groupby('Category').agg(
    Orders=('Order_ID', 'nunique'),
    Units=('Quantity', 'sum'),
    Sales=('Sales', 'sum'),
    Profit=('Profit', 'sum')
).reset_index()
cat_df['Margin'] = cat_df['Profit'] / cat_df['Sales']

row_curr = 11
for _, r in cat_df.iterrows():
    ws_summary.cell(row=row_curr, column=2, value=r['Category']).font = regular_font
    ws_summary.cell(row=row_curr, column=3, value=r['Orders']).number_format = '#,##0'
    ws_summary.cell(row=row_curr, column=4, value=r['Units']).number_format = '#,##0'
    ws_summary.cell(row=row_curr, column=5, value=r['Sales']).number_format = '$#,##0.00'
    ws_summary.cell(row=row_curr, column=6, value=r['Profit']).number_format = '$#,##0.00'
    ws_summary.cell(row=row_curr, column=7, value=r['Margin']).number_format = '0.00%'
    
    for c in range(2, 8):
        ws_summary.cell(row=row_curr, column=c).border = thin_border
        if c > 2:
            ws_summary.cell(row=row_curr, column=c).alignment = Alignment(horizontal="right")
    row_curr += 1

# Total Row
ws_summary.cell(row=row_curr, column=2, value="Total / Overall").font = bold_font
ws_summary.cell(row=row_curr, column=3, value=tot_orders).number_format = '#,##0'
ws_summary.cell(row=row_curr, column=4, value=df['Quantity'].sum()).number_format = '#,##0'
ws_summary.cell(row=row_curr, column=5, value=tot_sales).number_format = '$#,##0.00'
ws_summary.cell(row=row_curr, column=6, value=tot_profit).number_format = '$#,##0.00'
ws_summary.cell(row=row_curr, column=7, value=margin_pct).number_format = '0.00%'

for c in range(2, 8):
    cell = ws_summary.cell(row=row_curr, column=c)
    cell.font = bold_font
    cell.border = total_border
    if c > 2:
        cell.alignment = Alignment(horizontal="right")


# ---------------------------------------------------------------------------
# SHEET 2: RAW CLEANED DATA
# ---------------------------------------------------------------------------
ws_data = wb.create_sheet(title="Cleaned Data")
ws_data.views.sheetView[0].showGridLines = True

headers = list(df.columns)
for col_idx, h in enumerate(headers, start=1):
    cell = ws_data.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row_idx, row in df.iterrows():
    r_num = row_idx + 2
    for col_idx, val in enumerate(row, start=1):
        cell = ws_data.cell(row=r_num, column=col_idx, value=val)
        cell.font = regular_font
        col_name = headers[col_idx - 1]
        
        if col_name in ['Unit_Price', 'Sales', 'Profit']:
            cell.number_format = '$#,##0.00'
            cell.alignment = Alignment(horizontal="right")
        elif col_name == 'Discount':
            cell.number_format = '0.0%'
            cell.alignment = Alignment(horizontal="right")
        elif col_name in ['Quantity']:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right")
        elif col_name == 'Order_Date':
            cell.alignment = Alignment(horizontal="center")


# ---------------------------------------------------------------------------
# SHEET 3: CATEGORY & SUB-CATEGORY BREAKDOWN
# ---------------------------------------------------------------------------
ws_subcat = wb.create_sheet(title="Category Analysis")
ws_subcat.views.sheetView[0].showGridLines = True

ws_subcat["B2"] = "Category & Sub-Category Performance Analysis"
ws_subcat["B2"].font = title_font

subcat_df = df.groupby(['Category', 'Sub_Category']).agg(
    Orders=('Order_ID', 'nunique'),
    Units=('Quantity', 'sum'),
    Sales=('Sales', 'sum'),
    Profit=('Profit', 'sum')
).reset_index()
subcat_df['Margin'] = subcat_df['Profit'] / subcat_df['Sales']

sub_headers = ["Category", "Sub-Category", "Orders", "Units Sold", "Sales ($)", "Profit ($)", "Margin (%)"]
for col_idx, h in enumerate(sub_headers, start=2):
    cell = ws_subcat.cell(row=4, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center" if col_idx > 3 else "left")

row_curr = 5
for _, r in subcat_df.iterrows():
    ws_subcat.cell(row=row_curr, column=2, value=r['Category']).font = regular_font
    ws_subcat.cell(row=row_curr, column=3, value=r['Sub_Category']).font = regular_font
    ws_subcat.cell(row=row_curr, column=4, value=r['Orders']).number_format = '#,##0'
    ws_subcat.cell(row=row_curr, column=5, value=r['Units']).number_format = '#,##0'
    ws_subcat.cell(row=row_curr, column=6, value=r['Sales']).number_format = '$#,##0.00'
    ws_subcat.cell(row=row_curr, column=7, value=r['Profit']).number_format = '$#,##0.00'
    ws_subcat.cell(row=row_curr, column=8, value=r['Margin']).number_format = '0.00%'
    
    for c in range(2, 9):
        ws_subcat.cell(row=row_curr, column=c).border = thin_border
        if c >= 4:
            ws_subcat.cell(row=row_curr, column=c).alignment = Alignment(horizontal="right")
    row_curr += 1


# ---------------------------------------------------------------------------
# SHEET 4: REGIONAL PERFORMANCE
# ---------------------------------------------------------------------------
ws_reg = wb.create_sheet(title="Regional Performance")
ws_reg.views.sheetView[0].showGridLines = True

ws_reg["B2"] = "Geographic & Regional Performance Matrix"
ws_reg["B2"].font = title_font

reg_df = df.groupby('Region').agg(
    Orders=('Order_ID', 'nunique'),
    Units=('Quantity', 'sum'),
    Sales=('Sales', 'sum'),
    Profit=('Profit', 'sum')
).reset_index()
reg_df['Margin'] = reg_df['Profit'] / reg_df['Sales']

reg_headers = ["Region", "Orders", "Units Sold", "Sales ($)", "Profit ($)", "Margin (%)"]
for col_idx, h in enumerate(reg_headers, start=2):
    cell = ws_reg.cell(row=4, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left")

row_curr = 5
for _, r in reg_df.iterrows():
    ws_reg.cell(row=row_curr, column=2, value=r['Region']).font = regular_font
    ws_reg.cell(row=row_curr, column=3, value=r['Orders']).number_format = '#,##0'
    ws_reg.cell(row=row_curr, column=4, value=r['Units']).number_format = '#,##0'
    ws_reg.cell(row=row_curr, column=5, value=r['Sales']).number_format = '$#,##0.00'
    ws_reg.cell(row=row_curr, column=6, value=r['Profit']).number_format = '$#,##0.00'
    ws_reg.cell(row=row_curr, column=7, value=r['Margin']).number_format = '0.00%'
    
    for c in range(2, 8):
        ws_reg.cell(row=row_curr, column=c).border = thin_border
        if c >= 3:
            ws_reg.cell(row=row_curr, column=c).alignment = Alignment(horizontal="right")
    row_curr += 1


# Auto-fit column widths across all worksheets
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                val_str = str(cell.value)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(excel_path)
print(f"Professional Excel workbook saved to: {excel_path}")
